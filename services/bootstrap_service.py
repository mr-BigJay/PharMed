from openpyxl import load_workbook

from app_paths import resource_path
from services.db_session import SessionLocal

from models.category import Category
from models.center import Center
from models.health_house import HealthHouse
from models.item import Item


DRUG_CATEGORY_NAME = "دارو"
MEDICAL_CATEGORY_NAME = "تجهیزات پزشکی"


def bootstrap_reference_data(
    verbose=False
):
    db = SessionLocal()

    try:
        stats = {
            "categories": 0,
            "centers": 0,
            "health_houses": 0,
            "drugs": 0,
            "medical_items": 0,
        }

        categories = _ensure_categories(
            db
        )
        stats["categories"] = categories["added"]

        centers_file = resource_path(
            "data",
            "centers.xlsx"
        )
        health_houses_file = resource_path(
            "data",
            "health_houses.xlsx"
        )
        inventory_items_file = resource_path(
            "data",
            "inventory_items.xlsx"
        )

        if centers_file.exists():
            stats["centers"] = _import_centers(
                db,
                centers_file
            )

        if health_houses_file.exists():
            stats["health_houses"] = _import_health_houses(
                db,
                health_houses_file
            )

        if inventory_items_file.exists():
            drug_added, medical_added = _import_items(
                db,
                inventory_items_file,
                categories["drug"],
                categories["medical"]
            )
            stats["drugs"] = drug_added
            stats["medical_items"] = medical_added

        db.commit()

        if verbose:
            print(
                "Reference data bootstrap:",
                stats
            )

        return stats

    except Exception:
        db.rollback()
        raise

    finally:
        db.close()


def _ensure_categories(
    db
):
    added = 0

    drug_category = (
        db.query(Category)
        .filter_by(name=DRUG_CATEGORY_NAME)
        .first()
    )
    if not drug_category:
        drug_category = Category(
            name=DRUG_CATEGORY_NAME
        )
        db.add(
            drug_category
        )
        added += 1

    medical_category = (
        db.query(Category)
        .filter_by(name=MEDICAL_CATEGORY_NAME)
        .first()
    )
    if not medical_category:
        medical_category = Category(
            name=MEDICAL_CATEGORY_NAME
        )
        db.add(
            medical_category
        )
        added += 1

    db.flush()

    return {
        "added": added,
        "drug": drug_category,
        "medical": medical_category,
    }


def _import_centers(
    db,
    excel_file
):
    added = 0
    next_number = db.query(Center).count() + 1

    for row in _iter_sheet_rows(
        excel_file
    ):
        center_name = _clean_cell(
            row[0]
            if row
            else None
        )

        if not center_name:
            continue

        exists = (
            db.query(Center)
            .filter_by(name=center_name)
            .first()
        )
        if exists:
            continue

        center_code, next_number = _next_code(
            db,
            Center,
            "C",
            next_number
        )
        center = Center(
            code=center_code,
            name=center_name,
            is_active=True
        )
        db.add(
            center
        )
        added += 1

    db.flush()

    return added


def _import_health_houses(
    db,
    excel_file
):
    added = 0
    next_number = db.query(HealthHouse).count() + 1

    for row in _iter_sheet_rows(
        excel_file
    ):
        center_name = _clean_cell(
            row[0]
            if len(row) > 0
            else None
        )
        house_name = _clean_cell(
            row[1]
            if len(row) > 1
            else None
        )

        if not center_name or not house_name:
            continue

        center = (
            db.query(Center)
            .filter_by(name=center_name)
            .first()
        )
        if not center:
            continue

        exists = (
            db.query(HealthHouse)
            .filter_by(
                center_id=center.id,
                name=house_name
            )
            .first()
        )
        if exists:
            continue

        house_code, next_number = _next_code(
            db,
            HealthHouse,
            "H",
            next_number
        )
        health_house = HealthHouse(
            center_id=center.id,
            code=house_code,
            name=house_name,
            is_active=True
        )
        db.add(
            health_house
        )
        added += 1

    db.flush()

    return added


def _import_items(
    db,
    excel_file,
    drug_category,
    medical_category
):
    drug_added = _import_drugs(
        db,
        excel_file,
        drug_category
    )
    medical_added = _import_medical_items(
        db,
        excel_file,
        medical_category
    )

    db.flush()

    return (
        drug_added,
        medical_added
    )


def _import_drugs(
    db,
    excel_file,
    drug_category
):
    added = 0

    for row in _iter_sheet_rows(
        excel_file,
        sheet_name="دارو ها",
        skip_first_row=True
    ):
        item_name = _clean_cell(
            row[0]
            if len(row) > 0
            else None
        )
        item_form = _clean_cell(
            row[1]
            if len(row) > 1
            else None
        )

        if not item_name:
            continue

        exists = (
            db.query(Item)
            .filter_by(
                category_id=drug_category.id,
                item_name=item_name,
                item_form=item_form
            )
            .first()
        )
        if exists:
            continue

        item = Item(
            category_id=drug_category.id,
            item_name=item_name,
            item_form=item_form,
            unit="عدد",
            minimum_stock=0
        )
        db.add(
            item
        )
        added += 1

    return added


def _import_medical_items(
    db,
    excel_file,
    medical_category
):
    added = 0

    for row in _iter_sheet_rows(
        excel_file,
        sheet_name="تجهیزات پزشکی",
        skip_first_row=True
    ):
        item_name = _clean_cell(
            row[0]
            if len(row) > 0
            else None
        )

        if not item_name:
            continue

        exists = (
            db.query(Item)
            .filter_by(
                category_id=medical_category.id,
                item_name=item_name
            )
            .first()
        )
        if exists:
            continue

        item = Item(
            category_id=medical_category.id,
            item_name=item_name,
            item_form="",
            unit="عدد",
            minimum_stock=0
        )
        db.add(
            item
        )
        added += 1

    return added


def _iter_sheet_rows(
    excel_file,
    sheet_name=None,
    skip_first_row=False
):
    workbook = load_workbook(
        excel_file,
        read_only=True,
        data_only=True
    )

    try:
        sheet = (
            workbook[sheet_name]
            if sheet_name
            else workbook.active
        )
        rows = sheet.iter_rows(
            values_only=True
        )

        if skip_first_row:
            next(
                rows,
                None
            )

        yield from rows

    finally:
        workbook.close()


def _clean_cell(
    value
):
    if value is None:
        return ""

    return str(
        value
    ).strip()


def _next_code(
    db,
    model,
    prefix,
    start_number
):
    number = start_number

    while True:
        code = f"{prefix}{number:03}"
        exists = (
            db.query(model)
            .filter_by(code=code)
            .first()
        )

        if not exists:
            return (
                code,
                number + 1
            )

        number += 1
